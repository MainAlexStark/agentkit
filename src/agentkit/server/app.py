"""HTTP API агента и веб-чат.

Прогоны идут в фоне (RunManager), клиент следит за ними по SSE
(/api/sessions/{id}/events) или опросом (/api/sessions/{id}/poll) — второй
вариант для хостов, где прокси не держит долгие соединения.
"""

from __future__ import annotations

import asyncio
import html
import json
import re
import time
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, Literal

from fastapi import APIRouter, Depends, FastAPI, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from agentkit.attachments import UPLOADS, save_upload
from agentkit.events import EventType
from agentkit.mcp import MCPAuthRequired, MCPClient, MCPServerConfig
from agentkit.oauth import OAuthError, OAuthOptions
from agentkit.permissions import Mode
from agentkit.runs import Busy
from agentkit.scheduler import next_run, validate_cron
from agentkit.server.auth import require_admin
from agentkit.server.platform_api import register as register_platform
from agentkit.server.service import MCP_NAME, AgentService
from agentkit.store.base import ScheduleRec
from agentkit.tools.base import User

STATIC = Path(__file__).parent / "static"


class NewSession(BaseModel):
    title: str = ""
    model: str | None = None
    mode: str | None = None
    project_id: str | None = None


class PatchSession(BaseModel):
    title: str | None = None
    model: str | None = None
    mode: str | None = None


class SendMessage(BaseModel):
    text: str = ""
    files: list[dict[str, Any]] = []  # то, что вернул POST /files
    model: str | None = None
    mode: str | None = None


class Decide(BaseModel):
    approve: bool
    comment: str = ""
    args: dict[str, Any] | None = None


class Rewind(BaseModel):
    seq: int


class ProfileIn(BaseModel):
    name: str = ""  # полное имя
    call_me: str = ""  # как обращаться
    work: str = ""  # чем занимается
    instructions: str = ""  # пожелания к ответам — уходят в системный промпт


class MemoryIn(BaseModel):
    text: str
    scope: str = "user"  # user | global


class MemoryPatch(BaseModel):
    text: str


class ModelPrefs(BaseModel):
    favorites: list[str] | None = None
    only_favorites: bool | None = None
    default: str | None = None


class ModelPolicy(BaseModel):
    ids: list[str] | None = None  # None — снять ограничение (действует AGENT_MODELS_ALLOW)


class ScheduleIn(BaseModel):
    id: str | None = None  # есть — правка существующей задачи
    name: str
    cron: str
    prompt: str
    enabled: bool = True
    keep_session: bool = False


class SchedulePatch(BaseModel):
    enabled: bool


class MCPServerIn(BaseModel):
    name: str
    url: str
    auth: Literal["none", "token", "oauth"] = "none"
    token: str | None = None  # None — оставить прежний (в интерфейс токен не отдаётся)
    headers: dict[str, str] | None = None
    enabled: bool = True
    deferred: bool | None = None
    original_name: str | None = None  # правка существующего (в том числе переименование)


class MCPTestIn(BaseModel):
    url: str
    auth: Literal["none", "token", "oauth"] = "none"
    token: str | None = None
    headers: dict[str, str] | None = None
    name: str | None = None  # для правки: взять сохранённый токен


class MCPToggle(BaseModel):
    enabled: bool


class HeadlessRun(BaseModel):
    text: str
    session_id: str | None = None
    model: str | None = None
    mode: str | None = None


def build_app(svc: AgentService, *, extra_routers: list[APIRouter] | None = None,
              static_dir: Path | None = None) -> FastAPI:
    @asynccontextmanager
    async def lifespan(app: FastAPI):
        await svc.start()
        yield
        await svc.stop()

    app = FastAPI(title=svc.settings.name, lifespan=lifespan)
    if svc.settings.allowed_origins:
        app.add_middleware(CORSMiddleware, allow_origins=[o.strip() for o in svc.settings.allowed_origins.split(",")],
                           allow_methods=["*"], allow_headers=["*"])
    auth = svc.auth
    api = APIRouter(prefix="/api")

    async def own_session(session_id: str, user: User):
        s = await svc.store.get_session(session_id)
        if s is None or (s.user_id != user.id and not user.has_role("admin")):
            raise HTTPException(404, "нет такой сессии")
        return s

    @app.get("/healthz")
    async def healthz():
        return {"ok": True, "agent": svc.settings.name, "tools": len(svc.registry or [])}

    @api.get("/me")
    async def me(user: User = Depends(auth)):
        prof = await svc.store.kv_get("profile", user.id) or {}
        return {"id": user.id, "name": prof.get("name") or user.name, "call_me": prof.get("call_me", ""),
                "roles": user.roles}

    @api.get("/profile")
    async def profile(user: User = Depends(auth)):
        return {"name": user.name, "call_me": "", "work": "", "instructions": "",
                **(await svc.store.kv_get("profile", user.id) or {})}

    @api.put("/profile")
    async def save_profile(body: ProfileIn, user: User = Depends(auth)):
        data = {k: v.strip()[:4000] for k, v in body.model_dump().items()}
        await svc.store.kv_set("profile", user.id, data)
        return data

    @api.get("/plan")
    async def plan(user: User = Depends(auth)):
        now = time.localtime()
        month_start = time.mktime((now.tm_year, now.tm_mon, 1, 0, 0, 0, 0, 0, -1))
        usage = await svc.store.usage_summary(user_id=user.id, since=month_start)
        s = svc.settings
        info: dict[str, Any] = {"name": s.plan_name or "Без ограничений", "status": "active",
                                "limit_usd": s.monthly_budget_usd or None, "manage_url": s.billing_url or None,
                                "features": []}
        if svc.domain.plan is not None:
            res = svc.domain.plan(user, usage)
            if hasattr(res, "__await__"):
                res = await res
            info.update(res or {})
        return {**info, "usage": usage, "period_start": month_start}

    @api.get("/config")
    async def config(user: User = Depends(auth)):
        s = svc.settings
        prefs = await svc.model_prefs(user)
        policy = await svc.model_policy()
        favs = [m for m in prefs["favorites"] if svc.model_allowed(m, policy)]
        return {"name": s.name, "models": favs or [s.model_main],
                "default_model": prefs["default"] if svc.model_allowed(prefs["default"], policy) else s.model_main,
                "modes": [m.value for m in Mode if m != Mode.BYPASS or user.has_role("admin")],
                "default_mode": s.default_mode, "is_admin": user.has_role("admin"),
                "description": svc.domain.description or s.description,
                "suggestions": svc.domain.suggestions or [x.strip() for x in s.suggestions.split("|") if x.strip()],
                "titles": {t.name: t.title for t in svc.registry.all() if t.title},
                "help_url": s.help_url or None, "plan_name": s.plan_name or None}

    # ── сессии ───────────────────────────────────────────────────────
    async def project_names(user: User) -> dict[str, str]:
        if svc.projects is None:
            return {}
        return {p.id: p.name for p, _ in await svc.store.project_list_for(user.id)}

    @api.get("/sessions")
    async def list_sessions(project_id: str | None = None, user: User = Depends(auth)):
        items = await svc.store.list_sessions(user.id, limit=100, project_id=project_id)
        names = await project_names(user)
        return [{"id": s.id, "title": s.title or "Новый разговор", "updated_at": s.updated_at,
                 "project_id": s.project_id if s.project_id in names else None, "project": names.get(s.project_id or ""),
                 "running": await svc.runs.active_run(s.id) is not None} for s in items]

    @api.post("/sessions")
    async def create_session(body: NewSession, user: User = Depends(auth)):
        state = {k: v for k, v in {"model": await svc.check_model(body.model),
                                   "mode": _check_mode(body.mode, user)}.items() if v}
        if body.project_id:
            await app.state.project_access(body.project_id, user)  # участник проекта — хотя бы читатель
        s = await svc.store.create_session(user.id, title=body.title, agent=svc.agent.spec.name, state=state,
                                           project_id=body.project_id or None)
        return {"id": s.id}

    @api.get("/sessions/{session_id}")
    async def get_session(session_id: str, user: User = Depends(auth)):
        s = await own_session(session_id, user)
        msgs = await svc.store.list_messages(session_id)
        times = await svc.store.message_times(session_id)
        pending = await svc.store.list_approvals(session_id, status="pending")
        project = None
        if s.project_id and svc.projects is not None:
            p, role = await svc.store.project_role(s.project_id, user.id)
            project = {"id": p.id, "name": p.name, "role": role} if p and role else None
        return {
            "id": s.id, "title": s.title, "state": s.state, "project": project,
            "messages": [{"seq": seq, "ts": times.get(seq), **m.to_dict()} for seq, m in msgs],
            "approvals": [{**a.__dict__, "destructive": bool((t := svc.registry.get(a.tool_name)) and t.destructive),
                           "title": t.title if t else ""} for a in pending],
            "running_run": await svc.runs.active_run(session_id),
            "queued": await svc.store.inbox_count(session_id),
            "last_seq": await svc.store.last_event_seq(session_id),
            "runs": [r.__dict__ for r in await svc.store.list_runs(session_id, limit=20)],
        }

    @api.patch("/sessions/{session_id}")
    async def patch_session(session_id: str, body: PatchSession, user: User = Depends(auth)):
        s = await own_session(session_id, user)
        if body.title is not None:
            await svc.store.update_session(session_id, title=body.title)
        if body.model is not None or body.mode is not None:
            state = dict(s.state or {})
            if body.model is not None:
                state["model"] = await svc.check_model(body.model)
            if body.mode is not None:
                state["mode"] = _check_mode(body.mode, user)
            await svc.store.set_state(session_id, state)
        return {"ok": True}

    @api.delete("/sessions/{session_id}")
    async def delete_session(session_id: str, user: User = Depends(auth)):
        await own_session(session_id, user)
        if await svc.runs.active_run(session_id):
            raise HTTPException(409, "сначала остановите прогон")
        await svc.store.delete_session(session_id)
        return {"ok": True}

    @api.post("/sessions/{session_id}/messages")
    async def send(session_id: str, body: SendMessage, user: User = Depends(auth)):
        """Если агент занят — сообщение встаёт в очередь, и он подхватит его между шагами (queued: true)."""
        await own_session(session_id, user)
        mode = _check_mode(body.mode, user)
        model = await svc.check_model(body.model)
        content = _content(session_id, body)
        if not content:
            raise HTTPException(400, "пустое сообщение")
        return await svc.send(session_id, user, content, model=model, mode=mode)

    @api.post("/sessions/{session_id}/cancel")
    async def cancel(session_id: str, user: User = Depends(auth)):
        await own_session(session_id, user)
        return {"ok": await svc.runs.cancel(session_id)}

    @api.post("/sessions/{session_id}/rewind")
    async def rewind(session_id: str, body: Rewind, user: User = Depends(auth)):
        await own_session(session_id, user)
        if await svc.runs.active_run(session_id):
            raise HTTPException(409, "сначала остановите прогон")
        return await svc.agent.rewind(session_id, body.seq)

    @api.get("/sessions/{session_id}/events")
    async def events(session_id: str, after: int = 0, user: User = Depends(auth)):
        await own_session(session_id, user)

        async def gen():
            yield ": ok\n\n"
            async for ev in svc.runs.subscribe(session_id, after):
                yield f"id: {ev.get('seq', 0)}\nevent: {ev['type']}\ndata: {json.dumps(ev, ensure_ascii=False, default=str)}\n\n"
            yield "event: idle\ndata: {}\n\n"

        return StreamingResponse(gen(), media_type="text/event-stream",
                                 headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})

    @api.get("/sessions/{session_id}/poll")
    async def poll(session_id: str, after: int = 0, user: User = Depends(auth)):
        await own_session(session_id, user)
        evs = await svc.store.list_events(session_id=session_id, after_seq=after, limit=500)
        return {"events": [e.to_dict() for e in evs], "running": await svc.runs.active_run(session_id) is not None}

    # ── подтверждения ────────────────────────────────────────────────
    @api.post("/approvals/{approval_id}")
    async def decide(approval_id: str, body: Decide, user: User = Depends(auth)):
        rec = await svc.store.get_approval(approval_id)
        if rec is None:
            raise HTTPException(404, "нет такого подтверждения")
        await own_session(rec.session_id, user)
        if rec.status != "pending":
            raise HTTPException(409, f"уже решено: {rec.status}")
        try:
            run_id = await svc.runs.start(
                rec.session_id,
                lambda: svc.agent.resolve_approval(approval_id, user, approve=body.approve, comment=body.comment,
                                                   edited_args=body.args),
                cancel=svc.agent.cancel)
        except Busy:
            raise HTTPException(409, "агент ещё работает в этом разговоре") from None
        return {"run_id": run_id}

    # ── синхронный запуск для интеграций ─────────────────────────────
    @api.post("/run")
    async def headless(body: HeadlessRun, user: User = Depends(auth)):
        """Для других сервисов: отправить задачу и дождаться итога одним запросом."""
        sid = body.session_id
        if sid:
            await own_session(sid, user)
        else:
            sid = (await svc.store.create_session(user.id, title=body.text[:80])).id
        mode = _check_mode(body.mode, user)
        model = await svc.check_model(body.model)
        try:
            # Через RunManager, а не напрямую: прогон виден в веб-чате, его можно
            # остановить, и второй запрос в ту же сессию получит 409, а не гонку.
            run_id = await svc.runs.start(
                sid, lambda: svc.agent.run(sid, user, body.text, model=model, mode=mode, trigger="api"),
                cancel=svc.agent.cancel)
        except Busy:
            raise HTTPException(409, "в сессии уже идёт прогон") from None
        await svc.runs.wait(sid)
        run = await svc.store.get_run(run_id)
        evs = await svc.store.list_events(session_id=sid, limit=100_000)
        text = next((e.data.get("content") for e in reversed(evs)
                     if e.run_id == run_id and e.type == EventType.MESSAGE and not e.agent and e.data.get("content")), "")
        pending = await svc.store.list_approvals(sid, status="pending")
        return {"session_id": sid, "run_id": run_id, "status": run.status if run else "unknown", "text": text,
                "usage": run.usage if run else {}, "approvals": [a.__dict__ for a in pending]}

    # ── память ───────────────────────────────────────────────────────
    @api.get("/memory")
    async def memory_list(user: User = Depends(auth)):
        items = await svc.store.memory_list(["global", f"user:{user.id}"])
        return [m.__dict__ for m in items]

    @api.post("/memory")
    async def memory_add(body: MemoryIn, user: User = Depends(auth)):
        if body.scope == "global":
            require_admin(user)
        rec = await svc.store.memory_add("global" if body.scope == "global" else f"user:{user.id}", body.text,
                                         created_by=user.id)
        return rec.__dict__

    async def own_memory(memory_id: str, user: User) -> None:
        """Своё правит каждый, общее («для всех») — только админ."""
        items = {m.id: m for m in await svc.store.memory_list(["global", f"user:{user.id}"])}
        if memory_id not in items:
            raise HTTPException(404, "нет такой записи")
        if items[memory_id].scope != f"user:{user.id}":
            require_admin(user)

    @api.patch("/memory/{memory_id}")
    async def memory_edit(memory_id: str, body: MemoryPatch, user: User = Depends(auth)):
        await own_memory(memory_id, user)
        if not body.text.strip():
            raise HTTPException(400, "пустая запись — удалите её вместо этого")
        await svc.store.memory_update(memory_id, body.text.strip())
        return {"ok": True}

    @api.delete("/memory/{memory_id}")
    async def memory_delete(memory_id: str, user: User = Depends(auth)):
        await own_memory(memory_id, user)
        return {"ok": await svc.store.memory_delete(memory_id)}

    # ── инструменты, MCP, расход ─────────────────────────────────────
    @api.get("/tools")
    async def tools(user: User = Depends(auth)):
        custom = await svc.mcp_custom()
        return {
            "tools": [{"name": t.name, "title": t.title, "description": t.description.split("\n")[0], "read_only": t.read_only,
                       "destructive": t.destructive, "deferred": t.deferred, "group": t.group, "source": t.source}
                      for t in svc.registry.all()],
            "mcp": [_mcp_view(c, svc.mcp.status.get(c.name), custom.get(c.name), user.has_role("admin"),
                              c.name not in svc.mcp_file_names) for c in svc.mcp_all],
            "skills": [{"name": k.name, "description": k.description} for k in (svc.agent.skills.skills.values()
                                                                               if svc.agent.skills else [])],
        }

    # ── модели ───────────────────────────────────────────────────────
    @api.get("/models")
    async def models(refresh: bool = False, user: User = Depends(auth)):
        """Каталог моделей. Обычный пользователь видит только разрешённые; админ — все, с отметкой."""
        admin = user.has_role("admin")
        await svc.catalog.refresh(force=refresh and admin)
        policy = await svc.model_policy()
        items = []
        for m in svc.catalog.models():
            ok = svc.model_allowed(m.id, policy)
            if ok or admin:
                items.append({**m.to_dict(), "allowed": ok})
        prefs = await svc.model_prefs(user)
        out: dict[str, Any] = {"models": items, **prefs, "main": svc.settings.model_main,
                               "fetched_at": svc.catalog.fetched_at}
        if admin:
            out["policy"] = {"custom": policy is not None, "patterns": svc.settings.models_allow,
                             "errors": svc.catalog.errors}
        return out

    @api.put("/models/prefs")
    async def models_prefs(body: ModelPrefs, user: User = Depends(auth)):
        cur = await svc.store.kv_get("model_prefs", user.id) or {}
        if body.favorites is not None:
            cur["favorites"] = list(dict.fromkeys(body.favorites))[:200]
        if body.only_favorites is not None:
            cur["only_favorites"] = body.only_favorites
        if body.default is not None:
            cur["default"] = await svc.check_model(body.default) or ""
        await svc.store.kv_set("model_prefs", user.id, cur)
        return await svc.model_prefs(user)

    @api.put("/models/policy")
    async def models_policy(body: ModelPolicy, user: User = Depends(auth)):
        require_admin(user)
        if body.ids is None:
            await svc.store.kv_delete("config", "models_allowed")
        else:
            await svc.store.kv_set("config", "models_allowed", {"ids": sorted(set(body.ids))})
        return {"ok": True}

    @api.post("/reload")
    async def reload(user: User = Depends(auth)):
        require_admin(user)
        await svc.reload()
        return {"ok": True, "tools": len(svc.registry)}

    @api.get("/usage")
    async def usage(days: int = 30, user: User = Depends(auth)):
        since = time.time() - days * 86400
        return await svc.store.usage_summary(user_id=None if user.has_role("admin") else user.id, since=since)

    # ── расписание ───────────────────────────────────────────────────
    @api.get("/schedules")
    async def schedules(user: User = Depends(auth)):
        return [s.__dict__ for s in await svc.store.schedule_list(user_id=None if user.has_role("admin") else user.id)]

    async def own_schedule(schedule_id: str, user: User) -> ScheduleRec:
        s = {x.id: x for x in await svc.store.schedule_list()}.get(schedule_id)
        if s is None or (s.user_id != user.id and not user.has_role("admin")):
            raise HTTPException(404, "нет такого расписания")
        return s

    @api.get("/schedules/preview")
    async def schedule_preview(cron: str, user: User = Depends(auth)):
        """Ближайшие запуски — чтобы человек видел, что именно он настроил, до сохранения."""
        if not validate_cron(cron):
            return {"ok": False, "next": []}
        out, t = [], time.time()
        for _ in range(3):
            t = next_run(cron, svc.settings.tz, t)
            out.append(t)
        return {"ok": True, "next": out}

    @api.post("/schedules")
    async def add_schedule(body: ScheduleIn, user: User = Depends(auth)):
        if not body.name.strip() or not body.prompt.strip():
            raise HTTPException(400, "нужны название и задание для агента")
        if not validate_cron(body.cron):
            raise HTTPException(400, "неверное cron-выражение (5 полей: минута час день месяц день_недели)")
        from agentkit.messages import new_id

        old = await own_schedule(body.id, user) if body.id else None
        sid = old.session_id if old else None
        if body.keep_session and not sid:
            sid = (await svc.store.create_session(user.id, title=f"⏰ {body.name}")).id
        elif not body.keep_session:
            sid = None
        rec = ScheduleRec(id=old.id if old else new_id("sc_"), name=body.name.strip(), cron=body.cron.strip(),
                          prompt=body.prompt.strip(), user_id=old.user_id if old else user.id, enabled=body.enabled,
                          session_id=sid, last_run_at=old.last_run_at if old else None,
                          next_run_at=next_run(body.cron, svc.settings.tz))
        await svc.store.schedule_upsert(rec)
        return rec.__dict__

    @api.patch("/schedules/{schedule_id}")
    async def patch_schedule(schedule_id: str, body: SchedulePatch, user: User = Depends(auth)):
        s = await own_schedule(schedule_id, user)
        s.enabled = body.enabled
        # После паузы считаем от «сейчас»: иначе включённая задача сразу сработала бы за пропущенное.
        s.next_run_at = next_run(s.cron, svc.settings.tz) if body.enabled else s.next_run_at
        await svc.store.schedule_upsert(s)
        return s.__dict__

    @api.post("/schedules/{schedule_id}/run")
    async def run_schedule(schedule_id: str, user: User = Depends(auth)):
        s = await own_schedule(schedule_id, user)
        if not await svc.scheduler.run_now(s):
            raise HTTPException(409, "задача уже выполняется")
        return {"ok": True}

    @api.delete("/schedules/{schedule_id}")
    async def del_schedule(schedule_id: str, user: User = Depends(auth)):
        await own_schedule(schedule_id, user)
        await svc.store.schedule_delete(schedule_id)
        return {"ok": True}

    # ── файлы ────────────────────────────────────────────────────────
    def _workspace(session_id: str) -> Path:
        ws = svc.agent.workspace_for(session_id)
        if ws is None:
            raise HTTPException(400, "у агента нет рабочей папки (AGENT_ENABLE_FILES=false)")
        return ws

    def _content(session_id: str, body: SendMessage) -> Any:
        if not body.files:
            return body.text.strip()
        ws = _workspace(session_id)
        parts: list[dict[str, Any]] = [{"type": "text", "text": body.text}] if body.text.strip() else []
        for f in body.files:
            path = str(f.get("path", ""))
            target = (ws / path).resolve()
            # Ссылку присылает клиент — проверяем, что это действительно загруженный файл этой сессии.
            if not path.startswith(UPLOADS + "/") or ws.resolve() not in target.parents or not target.is_file():
                raise HTTPException(400, f"нет загруженного файла {path}")
            parts.append({k: f[k] for k in ("type", "path", "name", "kind", "size", "mime") if k in f} | {"type": "file"})
        return parts

    @api.post("/sessions/{session_id}/files")
    async def upload(session_id: str, file: UploadFile, user: User = Depends(auth)):
        await own_session(session_id, user)
        data = await file.read(svc.settings.max_upload_mb * 1024 * 1024 + 1)
        if len(data) > svc.settings.max_upload_mb * 1024 * 1024:
            raise HTTPException(413, f"файл больше {svc.settings.max_upload_mb} МБ")
        return save_upload(_workspace(session_id), file.filename or "file", data, file.content_type or "")

    @api.get("/sessions/{session_id}/files/preview")
    async def preview(session_id: str, path: str, user: User = Depends(auth)):
        """Excel и Word для панели предпросмотра: браузер их не показывает, отдаём таблицу или текст."""
        await own_session(session_id, user)
        ws = _workspace(session_id).resolve()
        p = (ws / path).resolve()
        if ws not in p.parents or not p.is_file():
            raise HTTPException(404, "нет файла")
        ext = p.suffix.lower()
        try:
            if ext == ".xlsx":
                import openpyxl

                wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
                sheets = []
                for ws_ in wb.worksheets[:10]:
                    rows = []
                    for row in ws_.iter_rows(values_only=True):
                        if len(rows) >= 500:
                            break
                        rows.append(["" if v is None else str(v) for v in row])
                    while rows and not any(rows[-1]):
                        rows.pop()
                    sheets.append({"name": ws_.title, "rows": rows})
                return {"kind": "sheets", "sheets": sheets}
            if ext == ".docx":
                import docx

                d = docx.Document(str(p))
                blocks: list[dict[str, Any]] = []
                for para in d.paragraphs:
                    style = (para.style.name or "").lower() if para.style is not None else ""
                    level = int(style.split()[-1]) if style.startswith("heading") and style.split()[-1].isdigit() else 0
                    if para.text.strip():
                        blocks.append({"type": "h" if level else "p", "level": level, "text": para.text})
                for t in d.tables:
                    blocks.append({"type": "table", "rows": [[c.text for c in r.cells] for r in t.rows][:300]})
                return {"kind": "doc", "blocks": blocks}
        except ImportError as e:
            raise HTTPException(501, f"для предпросмотра нужен пакет {e.name}") from None
        raise HTTPException(415, "для этого типа предпросмотра нет")

    @api.get("/sessions/{session_id}/files")
    async def download(session_id: str, path: str, user: User = Depends(auth)):
        await own_session(session_id, user)
        ws = _workspace(session_id).resolve()
        target = (ws / path).resolve()
        if ws not in target.parents or not target.is_file():
            raise HTTPException(404, "нет такого файла")
        return FileResponse(target, filename=target.name)

    # ── OAuth MCP-серверов ───────────────────────────────────────────
    @api.post("/mcp/{name}/connect")
    async def mcp_connect(name: str, user: User = Depends(auth)):
        require_admin(user)
        cfg = svc.mcp.config(name)
        if cfg is None or not cfg.url:
            raise HTTPException(404, f"нет HTTP MCP-сервера {name}")
        try:
            url = await svc.oauth.start(name, cfg.url, OAuthOptions.from_config(cfg.oauth), user_id=user.id)
        except OAuthError as e:
            raise HTTPException(400, str(e)) from None
        return {"url": url, "redirect_uri": svc.oauth.redirect_uri}

    @api.post("/mcp/servers")
    async def mcp_save(body: MCPServerIn, user: User = Depends(auth)):
        """Добавить или изменить свой MCP-сервер (из mcp.json — только правкой файла)."""
        require_admin(user)
        servers = await svc.mcp_custom()
        old_name = body.original_name or body.name
        if not MCP_NAME.match(body.name):
            raise HTTPException(400, "имя — латиница в нижнем регистре, цифры, «-» и «_», до 32 символов")
        if body.name in svc.mcp_file_names:
            raise HTTPException(409, f"сервер {body.name} уже задан в mcp.json — выберите другое имя")
        if body.name != old_name and body.name in servers:
            raise HTTPException(409, f"сервер {body.name} уже есть")
        if not body.original_name and body.name in servers:
            raise HTTPException(409, f"сервер {body.name} уже есть")
        prev = servers.pop(old_name, {}) if body.original_name else {}
        token = prev.get("token", "") if body.token is None else body.token.strip()
        if body.auth == "token" and not token:
            raise HTTPException(400, "укажите токен или выберите другой способ входа")
        servers[body.name] = {
            "url": _check_mcp_url(body.url), "auth": body.auth, "token": token if body.auth == "token" else "",
            "headers": prev.get("headers", {}) if body.headers is None else _check_headers(body.headers),
            "enabled": body.enabled, "deferred": body.deferred,
            "added_by": prev.get("added_by") or user.id, "created_at": prev.get("created_at") or time.time(),
        }
        if body.original_name and old_name != body.name:
            await _forget_oauth(old_name)
        await svc.save_mcp_custom(servers)
        st = svc.mcp.status.get(body.name)
        return {"ok": True, "status": st.__dict__ if st else None}

    @api.post("/mcp/test")
    async def mcp_test(body: MCPTestIn, user: User = Depends(auth)):
        """Проверить сервер до сохранения: отвечает ли, сколько инструментов, нужен ли вход."""
        require_admin(user)
        token = body.token
        if token is None and body.name:
            token = (await svc.mcp_custom()).get(body.name, {}).get("token", "")
        cfg = MCPServerConfig(name=body.name or "test", url=_check_mcp_url(body.url),
                              token=(token or None) if body.auth == "token" else None,
                              headers=_check_headers(body.headers or {}), timeout=20)
        client = MCPClient(cfg)
        try:
            raw = await asyncio.wait_for(client.list_tools(), 20)
        except MCPAuthRequired:
            return {"ok": False, "needs_auth": True,
                    "error": "сервер просит вход" + ("" if body.auth == "oauth" else " — выберите OAuth или токен")}
        except Exception as e:  # noqa: BLE001 — любой сбой показываем человеку как есть
            return {"ok": False, "needs_auth": False, "error": str(e) or type(e).__name__}
        finally:
            await client.close()
        return {"ok": True, "tools": [t.get("title") or t["name"] for t in raw][:50], "count": len(raw),
                "server": client.server_info.get("name") or "", "version": client.server_info.get("version") or ""}

    @api.patch("/mcp/servers/{name}")
    async def mcp_toggle(name: str, body: MCPToggle, user: User = Depends(auth)):
        require_admin(user)
        servers = await svc.mcp_custom()
        if name not in servers:
            raise HTTPException(404, "такого сервера нет среди добавленных в настройках")
        servers[name]["enabled"] = body.enabled
        await svc.save_mcp_custom(servers)
        return {"ok": True}

    @api.delete("/mcp/servers/{name}")
    async def mcp_delete(name: str, user: User = Depends(auth)):
        require_admin(user)
        servers = await svc.mcp_custom()
        if name not in servers:
            raise HTTPException(404, "такого сервера нет среди добавленных в настройках")
        servers.pop(name)
        await _forget_oauth(name)
        await svc.save_mcp_custom(servers)
        return {"ok": True}

    async def _forget_oauth(name: str) -> None:
        try:
            await svc.oauth.disconnect(name)
        except Exception:  # noqa: BLE001 — входа могло и не быть
            pass

    @api.delete("/mcp/{name}/connect")
    async def mcp_disconnect(name: str, user: User = Depends(auth)):
        require_admin(user)
        await svc.oauth.disconnect(name)
        await svc.reload()
        return {"ok": True}

    @app.get("/oauth/callback", response_class=HTMLResponse)
    async def oauth_callback(state: str = "", code: str = "", error: str = "", error_description: str = ""):
        # Без авторизации по токену: сюда возвращает браузер с сервера OAuth.
        # Защита — одноразовый state, выданный только администратору.
        if error:
            return _oauth_page(False, f"{error}: {error_description}")
        try:
            server = await svc.oauth.finish(state, code)
        except OAuthError as e:
            return _oauth_page(False, str(e))
        await svc.reload()
        return _oauth_page(True, f"{server} подключён")

    register_platform(api, app, svc, auth)
    app.include_router(api)
    for r in extra_routers or []:
        app.include_router(r)

    static = static_dir if static_dir and static_dir.is_dir() else STATIC

    @app.get("/")
    async def index():
        # Ссылки на app.js/app.css — с версией по содержимому файлов, а сама страница не кэшируется.
        # Иначе после обновления Chrome эвристически берёт старый app.js из кэша к новой
        # разметке (Safari перепроверяет чаще) — и интерфейс ломается только в одном браузере.
        text = (static / "index.html").read_text(encoding="utf-8")
        for name in ("app.js", "app.css", "viz.js"):
            f = static / name
            if f.is_file():
                st = f.stat()
                ver = f"{int(st.st_mtime)}-{st.st_size}"
                text = text.replace(f'"/static/{name}"', f'"/static/{name}?v={ver}"')
        return HTMLResponse(text, headers={"Cache-Control": "no-cache"})

    app.mount("/static", _RevalidatedStatic(directory=static), name="static")

    @app.exception_handler(ValueError)
    async def value_error(_, exc: ValueError):
        return JSONResponse({"detail": str(exc)}, status_code=400)

    return app


class _RevalidatedStatic(StaticFiles):
    """Статика всегда перепроверяется (ETag → 304): без этого браузер держит старый файл неделями."""

    def file_response(self, *args: Any, **kwargs: Any):
        resp = super().file_response(*args, **kwargs)
        resp.headers["Cache-Control"] = "no-cache"
        return resp


def _oauth_page(ok: bool, text: str) -> str:
    # Окно согласия открывалось из чата — сообщаем ему результат и закрываемся.
    payload = json.dumps({"type": "agentkit-oauth", "ok": ok, "text": text}, ensure_ascii=False)
    return f"""<!doctype html><meta charset="utf-8"><title>{'Готово' if ok else 'Ошибка'}</title>
<body style="font:16px system-ui;padding:40px;text-align:center">
<p>{'✓' if ok else '✕'} {html.escape(text)}</p><p style="color:#888">Окно можно закрыть.</p>
<script>try{{window.opener&&window.opener.postMessage({payload},"*");setTimeout(()=>window.close(),1200)}}catch(e){{}}</script>"""


def _check_mcp_url(url: str) -> str:
    url = url.strip()
    if not re.match(r"^https?://[^\s/]+", url):
        raise HTTPException(400, "адрес сервера — http(s)://…, например https://example.com/mcp")
    return url


def _check_headers(headers: dict[str, str]) -> dict[str, str]:
    out = {}
    for k, v in headers.items():
        k, v = k.strip(), str(v).strip()
        if not k:
            continue
        if not re.match(r"^[A-Za-z0-9-]+$", k) or not v.isascii():
            raise HTTPException(400, f"заголовок «{k}»: имя — латиница, цифры и «-», значение — без кириллицы")
        out[k] = v
    return out


def _mcp_view(c: Any, st: Any, d: dict[str, Any] | None, admin: bool, custom: bool) -> dict[str, Any]:
    """Сервер для списка в настройках. Токены не отдаются никогда — только «задан/не задан»."""
    view: dict[str, Any] = {
        "name": c.name, "enabled": c.enabled, "custom": custom,
        "ok": bool(st and st.ok), "tools": st.tools if st else 0, "error": st.error if st else "",
        "needs_auth": bool(st and st.needs_auth), "oauth": bool(st and st.oauth),
        "server": (st.server_info.get("name") or "") if st else "", "version": (st.server_info.get("version") or "") if st else "",
        "transport": "http" if c.url else "stdio",
    }
    if admin:  # адрес — только админу: в нём бывают внутренние хосты
        view["address"] = c.url or c.command or ""
        if custom and d:
            view.update(auth=d.get("auth", "none"), has_token=bool(d.get("token")), header_names=sorted(d.get("headers") or {}),
                        deferred=d.get("deferred"), added_by=d.get("added_by", ""))
    return view


def _check_mode(mode: str | None, user: User) -> str | None:
    if mode is None or mode == "":
        return None
    try:
        m = Mode(mode)
    except ValueError:
        raise HTTPException(400, f"неизвестный режим {mode}") from None
    if m == Mode.BYPASS and not user.has_role("admin"):
        raise HTTPException(403, "режим bypass — только для admin")
    return m.value

